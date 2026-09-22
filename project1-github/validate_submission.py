"""Validate the clean Final Project submission without modifying experiment data."""

from __future__ import annotations

import ast
import json
from importlib.metadata import version
from pathlib import Path

import pandas as pd
from PIL import Image
from docx import Document
from pptx import Presentation


ROOT = Path(__file__).resolve().parent
checks = {}
required = ["README.md", "requirements.txt", "dataset_description.md", "Project_Report.docx",
            "Final_Project_Presentation.pptx", "source_code", "results", "figures", "weights"]
checks["required_deliverables_exist"] = all((ROOT / item).exists() for item in required)

expected_folders = {"classification", "segmentation", "training", "evaluation", "utils"}
checks["source_structure_complete"] = expected_folders.issubset({path.name for path in (ROOT / "source_code").iterdir() if path.is_dir()})
python_files = list(ROOT.rglob("*.py"))
for path in python_files: ast.parse(path.read_text(encoding="utf-8"), filename=str(path))
checks["all_python_syntax_valid"] = bool(python_files)
checks["no_cache_or_notebook"] = not list(ROOT.rglob("__pycache__")) and not list(ROOT.rglob("*.pyc")) and not list(ROOT.rglob("*.ipynb"))

weights = list((ROOT / "weights").glob("*.pth"))
checks["seven_required_weights"] = len(weights) == 7 and all(path.stat().st_size > 1_000_000 for path in weights)
table = pd.read_csv(ROOT / "results/classification_model_comparison.csv")
expected_models = {"ResNet50", "VGG16", "VGG19", "DenseNet121", "EfficientNetB0", "CNN"}
checks["six_classifiers_reported"] = set(table.Model) == expected_models
checks["metrics_nonempty_no_nan"] = not table[["Accuracy","Macro Precision","Macro Recall","Macro F1","Weighted F1"]].isna().any().any()
checks["canonical_mps_results"] = set(table.Split) == {"canonical"} and set(table["Inference Device"]) == {"mps"} and set(table["Validation Samples"]) == {2048}

seg = json.load((ROOT / "results/segmentation_metrics.json").open())
checks["segmentation_metrics_complete"] = all(key in seg and 0 <= seg[key] for key in ["loss","dice","iou","precision","recall"])
manifest = pd.read_csv(ROOT / "dataset_manifest/manifest.csv")
checks["dataset_manifest_complete"] = len(manifest) == 10015 and manifest.image_file.astype(bool).all() and manifest.mask_file.astype(bool).all()
description = (ROOT / "dataset_description.md").read_text(encoding="utf-8")
checks["official_dataset_link_present"] = "10.7910/DVN/DBW86T" in description

images = list((ROOT / "figures").glob("*.png")) + list((ROOT / "results/confusion_matrices").glob("*.png"))
images_valid = bool(images)
for path in images:
    try:
        with Image.open(path) as image: image.verify()
    except Exception: images_valid = False
checks["all_images_open"] = images_valid

presentation = Presentation(ROOT / "Final_Project_Presentation.pptx")
checks["presentation_15_slides"] = len(presentation.slides) == 15
checks["presentation_contains_images"] = sum(1 for slide in presentation.slides for shape in slide.shapes if shape.shape_type == 13) >= 7
document = Document(ROOT / "Project_Report.docx")
document_text = "\n".join(paragraph.text for paragraph in document.paragraphs)
required_sections = ["Abstract","Introduction","Related Work","Dataset","Methods","Classification","Segmentation","Experiments","Results","Discussion","Limitations","Future Work","References"]
checks["report_sections_complete"] = all(section in document_text for section in required_sections)
checks["report_contains_images"] = len(document.inline_shapes) >= 3
best = table.sort_values("Accuracy", ascending=False).iloc[0]
checks["report_numbers_match_results"] = best.Model in document_text and f"{best.Accuracy:.2%}" in document_text and f"{best['Macro F1']:.4f}" in document_text and f"{seg['dice']:.4f}" in document_text

readme = (ROOT / "README.md").read_text(encoding="utf-8")
checks["readme_sections_complete"] = all(f"## {heading}" in readme for heading in ["Project Overview","Environment","Installation","Dataset","Training","Evaluation","Segmentation","Results","Folder Structure"])
checks["missing_cnn_history_not_fabricated"] = "no training curve is claimed" in readme and not (ROOT / "results/cnn_history.json").exists()

requirements = {}
for line in (ROOT / "requirements.txt").read_text().splitlines():
    if "==" in line:
        package, pinned = line.split("==", 1); requirements[package] = pinned
checks["requirements_match_working_environment"] = all(version(package) == pinned for package, pinned in requirements.items())

checks = {name: bool(value) for name, value in checks.items()}
status = "PASS" if all(checks.values()) else "FAIL"
summary = {"status": status, "passed": sum(checks.values()), "total": len(checks), "checks": checks}
(ROOT / "submission_validation.json").write_text(json.dumps(summary, indent=2), encoding="utf-8")
print(json.dumps(summary, indent=2))
raise SystemExit(0 if status == "PASS" else 1)
